import sys
import os
import re
import importlib
import html
from xml.sax.saxutils import unescape

import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from reportlab.platypus import KeepTogether, Table, Paragraph, Image, HRFlowable, Spacer, CondPageBreak, PageBreak
from reportlab.graphics.shapes import Drawing, String, Group, Circle, Rect, Polygon, Line

# Helper: set background color on a table cell
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

# Helper: set cell padding/margins (in dxa: 20 dxa = 1 pt)
def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

# Helper: set table borders
def set_table_borders(table, color='E5E7EB', sz='4', val='single'):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="none"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def set_data_table_borders(table, color='D1D5DB', sz='4', val='single'):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def set_box_borders(table, color='1F2937', sz='16', val='single', left_only=True):
    tblPr = table._tbl.tblPr
    if left_only:
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
                <w:insideH w:val="none"/>
                <w:insideV w:val="none"/>
            </w:tblBorders>
        ''')
    else:
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideH w:val="none"/>
                <w:insideV w:val="none"/>
            </w:tblBorders>
        ''')
    tblPr.append(borders)

def set_no_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/>
            <w:right w:val="none"/><w:insideH w:val="none"/><w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

# Formatting parser for ReportLab XML/HTML inline strings
HTML_ENTITIES = {
    '&bull;': '• ',
    '&nbsp;': ' ',
    '&mdash;': '—',
    '&ndash;': '–',
    '&deg;': '°',
    '&plusmn;': '±',
    '&times;': '×',
    '&le;': '≤',
    '&ge;': '≥',
    '&ne;': '≠',
    '&approx;': '≈',
    '&infin;': '∞',
    '&rarr;': '→',
    '&larr;': '←',
    '&harr;': '↔',
    '&alpha;': 'α',
    '&beta;': 'β',
    '&gamma;': 'γ',
    '&delta;': 'δ',
    '&mu;': 'µ',
    '&pi;': 'π',
    '&ang;': '∠',
    '&perp;': '⊥',
    '&amp;': '&',
    '&lt;': '<',
    '&gt;': '>',
    '&quot;': '"',
    '&apos;': "'",
    '&#39;': "'",
}

def clean_html_entities(text):
    for ent, val in HTML_ENTITIES.items():
        text = text.replace(ent, val)
    return html.unescape(text)

def add_formatted_runs(paragraph, text, base_font_size=10.5, base_color='111827', default_bold=False, default_italic=False):
    text = str(text)
    # Clean up linebreaks within paragraph text
    text = text.replace('<br/>', '\n').replace('<br>', '\n').replace('<br />', '\n')
    
    tokens = re.split(r'(</?[bi]>|</?sub>|</?sup>|</?super>|<font[^>]*>|</font>)', text)
    is_bold = default_bold
    is_italic = default_italic
    is_sub = False
    is_sup = False
    
    for tok in tokens:
        if not tok:
            continue
        tok_lower = tok.lower()
        if tok_lower == '<b>':
            is_bold = True
        elif tok_lower == '</b>':
            is_bold = False
        elif tok_lower == '<i>':
            is_italic = True
        elif tok_lower == '</i>':
            is_italic = False
        elif tok_lower == '<sub>':
            is_sub = True
        elif tok_lower == '</sub>':
            is_sub = False
        elif tok_lower in ('<sup>', '<super>'):
            is_sup = True
        elif tok_lower in ('</sup>', '</super>'):
            is_sup = False
        elif tok_lower.startswith('<font') or tok_lower == '</font>':
            continue
        else:
            clean_str = clean_html_entities(tok)
            # Check if there are newline characters
            sub_lines = clean_str.split('\n')
            for line_idx, line in enumerate(sub_lines):
                if line_idx > 0:
                    # Add line break
                    paragraph.add_run('\n')
                if line:
                    run = paragraph.add_run(line)
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(base_font_size)
                    run.bold = is_bold
                    run.italic = is_italic
                    if is_sub:
                        run.font.subscript = True
                    if is_sup:
                        run.font.superscript = True
                    if base_color:
                        r = int(base_color[0:2], 16)
                        g = int(base_color[2:4], 16)
                        b = int(base_color[4:6], 16)
                        run.font.color.rgb = RGBColor(r, g, b)

def extract_heading_info(tbl):
    cells = tbl._cellvalues
    badge_text = ''
    heading_text = ''
    c0 = cells[0][0]
    if isinstance(c0, Drawing):
        for shape in c0.contents:
            if isinstance(shape, String):
                badge_text = shape.text
            elif isinstance(shape, Group):
                for sub in shape.contents:
                    if isinstance(sub, String):
                        badge_text = sub.text
    elif isinstance(c0, Paragraph):
        badge_text = c0.text
    elif isinstance(c0, str):
        badge_text = c0
        
    c1 = cells[0][1]
    if isinstance(c1, Paragraph):
        heading_text = c1.text
    elif isinstance(c1, str):
        heading_text = c1
    return badge_text, heading_text

def resolve_image_path(img_path, chapter_dir):
    if os.path.isabs(img_path) and os.path.exists(img_path):
        return img_path
    candidate1 = os.path.join(chapter_dir, 'assets', os.path.basename(img_path))
    if os.path.exists(candidate1):
        return candidate1
    candidate2 = os.path.join(chapter_dir, os.path.basename(img_path))
    if os.path.exists(candidate2):
        return candidate2
    return img_path

def extract_images_and_captions_deep(item, chapter_dir):
    """
    Recursively inspect item and find all images and their associated subcaptions/captions.
    Returns a list of dicts: [{'img_path': ..., 'width_cm': ..., 'caption': ...}]
    """
    results = []
    
    if isinstance(item, Image):
        path = getattr(item, 'filename', getattr(item, '_file', ''))
        path = resolve_image_path(path, chapter_dir)
        w_cm = getattr(item, '_width', 400) / 28.3465
        results.append({'img_path': path, 'width_cm': w_cm, 'caption': ''})
        return results
        
    if isinstance(item, list):
        # Could be [Table, Paragraph] (framed image + label)
        sub_img = None
        sub_cap = ''
        for sub in item:
            if isinstance(sub, Paragraph):
                sub_cap = sub.text
            elif isinstance(sub, Image):
                sub_img = sub
            elif isinstance(sub, Table):
                nested_res = extract_images_and_captions_deep(sub, chapter_dir)
                if nested_res:
                    results.extend(nested_res)
        if sub_img:
            path = getattr(sub_img, 'filename', getattr(sub_img, '_file', ''))
            path = resolve_image_path(path, chapter_dir)
            w_cm = getattr(sub_img, '_width', 200) / 28.3465
            results.append({'img_path': path, 'width_cm': w_cm, 'caption': sub_cap})
        elif results and sub_cap:
            for r in results:
                if not r['caption']:
                    r['caption'] = sub_cap
        return results

    if isinstance(item, Table):
        # Check if table represents a grid of images
        for row in item._cellvalues:
            for cell in row:
                cell_res = extract_images_and_captions_deep(cell, chapter_dir)
                results.extend(cell_res)
        return results

    if isinstance(item, KeepTogether):
        main_caption = ''
        for sub in item._content:
            if isinstance(sub, Paragraph) and ('Caption' in sub.style.name or 'Figure' in sub.text or 'Fig.' in sub.text):
                main_caption = sub.text
            else:
                sub_res = extract_images_and_captions_deep(sub, chapter_dir)
                results.extend(sub_res)
        if main_caption and results:
            results[0]['main_caption'] = main_caption
        return results

    return results

def build_docx_for_chapter(ch_num, chapter_title, mod_name, chapter_dir, out_docx_path):
    print(f"\nProcessing Chapter {ch_num}: {chapter_title}...")
    abs_dir = os.path.abspath(chapter_dir)
    if abs_dir not in sys.path:
        sys.path.insert(0, abs_dir)
        
    repo_root = os.path.abspath('.')
    if repo_root not in sys.path:
        sys.path.insert(0, repo_root)
        
    mod = importlib.import_module(mod_name)
    story = getattr(mod, 'story', [])
    
    doc = docx.Document()
    
    # Page setup matching neet_template.py: A4, 1.5cm left/right, 1.4cm top/bottom
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    
    # Title tracking
    title_emitted = False
    
    def emit_image_block(img_info_list, main_caption=''):
        if not img_info_list:
            return
            
        if len(img_info_list) == 1:
            info = img_info_list[0]
            img_path = info['img_path']
            w_cm = min(info['width_cm'], 16.0)
            
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(6)
            p_img.paragraph_format.space_after = Pt(2)
            p_img.paragraph_format.keep_with_next = True
            
            if os.path.exists(img_path):
                r_img = p_img.add_run()
                r_img.add_picture(img_path, width=Cm(w_cm))
            
            cap = info.get('caption') or main_caption
            if cap:
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(2)
                p_cap.paragraph_format.space_after = Pt(8)
                add_formatted_runs(p_cap, cap, base_font_size=9.0, base_color='4B5563', default_italic=True)
                
        else:
            # Multi-image row (2, 3, 4 images)
            n_cols = len(img_info_list)
            if n_cols > 4:
                # Break into rows of 3 or 2
                pass
            col_w_cm = 18.0 / n_cols
            side_tbl = doc.add_table(rows=1, cols=n_cols)
            side_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            side_tbl.autofit = False
            for ci in range(n_cols):
                side_tbl.columns[ci].width = Cm(col_w_cm)
                
            for ci, info in enumerate(img_info_list):
                c = side_tbl.rows[0].cells[ci]
                c.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
                set_cell_margins(c, top=60, bottom=60, left=40, right=40)
                
                img_path = info['img_path']
                p_c = c.paragraphs[0]
                p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_c.paragraph_format.space_before = Pt(0)
                p_c.paragraph_format.space_after = Pt(2)
                p_c.paragraph_format.keep_with_next = True
                
                if os.path.exists(img_path):
                    r_s = p_c.add_run()
                    # Fit inside column with padding
                    r_s.add_picture(img_path, width=Cm(min(info['width_cm'], col_w_cm - 0.4)))
                    
                sub_cap = info.get('caption', '')
                if sub_cap:
                    p_sub = c.add_paragraph()
                    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_sub.paragraph_format.space_before = Pt(2)
                    p_sub.paragraph_format.space_after = Pt(4)
                    add_formatted_runs(p_sub, sub_cap, base_font_size=8.5, base_color='4B5563', default_italic=True)
                    
            set_no_borders(side_tbl)
            
            if main_caption:
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(4)
                p_cap.paragraph_format.space_after = Pt(8)
                add_formatted_runs(p_cap, main_caption, base_font_size=9.0, base_color='4B5563', default_italic=True)
            else:
                p_sp = doc.add_paragraph()
                p_sp.paragraph_format.space_before = Pt(0)
                p_sp.paragraph_format.space_after = Pt(6)

    def process_story_element(el):
        nonlocal title_emitted
        
        # 1. Title Block detection
        if not title_emitted:
            is_title_tbl = False
            title_str = chapter_title
            if isinstance(el, Table) and len(el._cellvalues) == 1 and len(el._cellvalues[0]) >= 2:
                c0 = el._cellvalues[0][0]
                c1 = el._cellvalues[0][1]
                if isinstance(c0, Drawing) and isinstance(c1, Paragraph) and 'Title' in c1.style.name:
                    is_title_tbl = True
                    title_str = c1.text
            elif isinstance(el, Paragraph) and 'Title' in el.style.name:
                is_title_tbl = True
                title_str = el.text

            if is_title_tbl:
                p_chap = doc.add_paragraph()
                p_chap.paragraph_format.space_before = Pt(0)
                p_chap.paragraph_format.space_after = Pt(2)
                r_chap = p_chap.add_run(f"CHAPTER {ch_num}")
                r_chap.font.name = "Times New Roman"
                r_chap.font.size = Pt(11)
                r_chap.font.bold = True
                r_chap.font.color.rgb = RGBColor(107, 114, 128)
                
                p_title = doc.add_paragraph()
                p_title.paragraph_format.space_before = Pt(0)
                p_title.paragraph_format.space_after = Pt(6)
                r_title = p_title.add_run(clean_html_entities(title_str))
                r_title.font.name = "Times New Roman"
                r_title.font.size = Pt(24)
                r_title.font.bold = True
                r_title.font.color.rgb = RGBColor(17, 24, 39)
                title_emitted = True
                return

        # 2. HRFlowable
        if isinstance(el, HRFlowable):
            p_hr = doc.add_paragraph()
            p_hr.paragraph_format.space_before = Pt(0)
            p_hr.paragraph_format.space_after = Pt(10)
            p_hr_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="1" w:color="1F2937"/></w:pBdr>')
            p_hr._p.get_or_add_pPr().append(p_hr_border)
            return

        # 3. Spacers & CondPageBreaks & PageBreak
        if isinstance(el, Spacer) or isinstance(el, CondPageBreak):
            return
            
        if isinstance(el, PageBreak):
            doc.add_page_break()
            return

        # 4. KeepTogether structures
        if isinstance(el, KeepTogether):
            content = el._content
            
            # Check if this KeepTogether is a Heading banner
            is_heading = any(isinstance(sub, Table) and len(sub._cellvalues) == 1 and len(sub._cellvalues[0]) >= 2 and isinstance(sub._cellvalues[0][0], Drawing) for sub in content)
            if is_heading:
                tbl = [sub for sub in content if isinstance(sub, Table)][0]
                badge_text, heading_text = extract_heading_info(tbl)
                
                # Determine heading level
                if badge_text in [str(ch_num), f'{ch_num}.0', 'Summary', 'Exercises', 'E', 'Appendix', 'Recap'] or not ('.' in badge_text):
                    level = 1
                elif '.' in badge_text:
                    dots = badge_text.count('.')
                    if dots == 1:
                        level = 1
                    elif dots == 2:
                        level = 2
                    else:
                        level = 3
                else:
                    level = 2
                    
                banner = doc.add_table(rows=1, cols=2)
                banner.alignment = WD_TABLE_ALIGNMENT.LEFT
                banner.autofit = False
                
                badge_w = 1.3 if len(badge_text) <= 5 else 2.4
                banner.columns[0].width = Cm(badge_w)
                banner.columns[1].width = Cm(18.0 - badge_w)
                
                c0 = banner.rows[0].cells[0]
                c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_background(c0, '1F2937')
                set_cell_margins(c0, top=60, bottom=60, left=80, right=80)
                p0 = c0.paragraphs[0]
                p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p0.paragraph_format.space_before = Pt(0)
                p0.paragraph_format.space_after = Pt(0)
                r0 = p0.add_run(badge_text)
                r0.font.name = 'Times New Roman'
                r0.font.size = Pt(8.5 if len(badge_text) > 4 else 9.5)
                r0.font.bold = True
                r0.font.color.rgb = RGBColor(255, 255, 255)
                
                c1 = banner.rows[0].cells[1]
                c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(c1, top=60, bottom=60, left=120, right=80)
                p1 = c1.paragraphs[0]
                p1.paragraph_format.space_before = Pt(0)
                p1.paragraph_format.space_after = Pt(0)
                font_sz = {1: 13.5, 2: 11.5, 3: 10.5}.get(level, 11.0)
                add_formatted_runs(p1, heading_text, base_font_size=font_sz, base_color='111827', default_bold=True)
                
                tblPr = banner._tbl.tblPr
                if level == 1:
                    borders = parse_xml(f'''
                        <w:tblBorders {nsdecls("w")}>
                            <w:top w:val="none"/><w:left w:val="none"/>
                            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="1F2937"/>
                            <w:right w:val="none"/><w:insideH w:val="none"/><w:insideV w:val="none"/>
                        </w:tblBorders>
                    ''')
                else:
                    borders = parse_xml(f'''
                        <w:tblBorders {nsdecls("w")}>
                            <w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/>
                            <w:right w:val="none"/><w:insideH w:val="none"/><w:insideV w:val="none"/>
                        </w:tblBorders>
                    ''')
                tblPr.append(borders)
                
                p_spacer = doc.add_paragraph()
                p_spacer.paragraph_format.space_before = Pt(0)
                p_spacer.paragraph_format.space_after = Pt(4)
                return

            # Check if KeepTogether contains Image(s)
            img_results = extract_images_and_captions_deep(el, abs_dir)
            if img_results:
                main_cap = img_results[0].get('main_caption', '')
                emit_image_block(img_results, main_cap)
                return

            # Otherwise, unwrap and process sub-elements
            for sub in content:
                process_story_element(sub)
            return

        # 5. Paragraphs
        if isinstance(el, Paragraph):
            style_name = el.style.name
            p = doc.add_paragraph()
            
            if 'Bullet1' in style_name:
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.15
                add_formatted_runs(p, el.text, base_font_size=10.5, base_color='111827')
            elif 'Bullet2' in style_name:
                p.paragraph_format.left_indent = Cm(0.8)
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.15
                add_formatted_runs(p, el.text, base_font_size=10.5, base_color='111827')
            elif 'Bullet3' in style_name:
                p.paragraph_format.left_indent = Cm(1.1)
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(2.5)
                p.paragraph_format.line_spacing = 1.15
                add_formatted_runs(p, el.text, base_font_size=10.2, base_color='111827')
            elif style_name == 'Body' or 'Body-' in style_name:
                p.paragraph_format.left_indent = Cm(0)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.15
                add_formatted_runs(p, el.text, base_font_size=10.5, base_color='111827')
            elif 'NoteBox' in style_name:
                p.paragraph_format.left_indent = Cm(0.3)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.15
                add_formatted_runs(p, el.text, base_font_size=10.0, base_color='1F2937', default_italic=True)
            elif 'Caption' in style_name:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(8)
                add_formatted_runs(p, el.text, base_font_size=9.0, base_color='4B5563', default_italic=True)
            else:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(3)
                add_formatted_runs(p, el.text, base_font_size=10.5, base_color='111827')
            return

        # 6. Tables
        if isinstance(el, Table):
            cells = el._cellvalues
            r_count = len(cells)
            c_count = len(cells[0]) if r_count > 0 else 0
            first_c = cells[0][0] if r_count > 0 and c_count > 0 else None
            
            # 6a. Check if this table contains figures/images
            img_results = extract_images_and_captions_deep(el, abs_dir)
            if img_results:
                emit_image_block(img_results)
                return

            # 6b. Keyterm (1x2 table with Drawing in 0,0 and Paragraph in 0,1)
            if r_count == 1 and c_count == 2 and isinstance(first_c, Drawing):
                p_kt = doc.add_paragraph()
                p_kt.paragraph_format.left_indent = Cm(0.5)
                p_kt.paragraph_format.space_before = Pt(1)
                p_kt.paragraph_format.space_after = Pt(3)
                r_bullet = p_kt.add_run("•  ")
                r_bullet.font.name = "Times New Roman"
                r_bullet.font.size = Pt(10.5)
                r_bullet.font.bold = True
                r_bullet.font.color.rgb = RGBColor(31, 41, 55)
                kt_text = cells[0][1].text if isinstance(cells[0][1], Paragraph) else str(cells[0][1])
                add_formatted_runs(p_kt, kt_text, base_font_size=10.5, base_color='111827')
                return

            # 6c. Note / Memory Aid (1x1 table containing Table)
            if r_count == 1 and c_count == 1 and isinstance(first_c, Table):
                nested = first_c
                box_tbl = doc.add_table(rows=1, cols=1)
                box_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                box_tbl.autofit = False
                box_tbl.columns[0].width = Cm(18.0)
                
                box_cell = box_tbl.rows[0].cells[0]
                set_cell_background(box_cell, 'F9FAFB')
                set_cell_margins(box_cell, top=100, bottom=100, left=150, right=150)
                
                para_obj = None
                for nr in nested._cellvalues:
                    for nc in nr:
                        if isinstance(nc, Paragraph):
                            para_obj = nc
                            break
                            
                p_box = box_cell.paragraphs[0]
                p_box.paragraph_format.space_before = Pt(0)
                p_box.paragraph_format.space_after = Pt(0)
                p_box.paragraph_format.line_spacing = 1.15
                
                box_text = para_obj.text if para_obj else ''
                is_mem = 'MEMORY AID' in box_text
                border_color = '4B5563' if is_mem else '1F2937'
                border_val = 'dashed' if is_mem else 'single'
                
                set_box_borders(box_tbl, color=border_color, sz='16', val=border_val, left_only=True)
                
                if para_obj:
                    add_formatted_runs(p_box, box_text, base_font_size=10.0, base_color='1F2937')
                    
                p_sp = doc.add_paragraph()
                p_sp.paragraph_format.space_before = Pt(0)
                p_sp.paragraph_format.space_after = Pt(6)
                return

            # 6d. Process flow (Table with Drawings in col 0)
            if r_count >= 2 and c_count == 2 and isinstance(first_c, Drawing):
                flow_tbl = doc.add_table(rows=r_count, cols=2)
                flow_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                flow_tbl.autofit = False
                flow_tbl.columns[0].width = Cm(1.2)
                flow_tbl.columns[1].width = Cm(16.8)
                
                for s_idx in range(r_count):
                    step_c0 = flow_tbl.rows[s_idx].cells[0]
                    step_c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    set_cell_background(step_c0, '1F2937')
                    set_cell_margins(step_c0, top=60, bottom=60, left=60, right=60)
                    sp0 = step_c0.paragraphs[0]
                    sp0.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    sr0 = sp0.add_run(f"{s_idx + 1}")
                    sr0.font.name = "Times New Roman"
                    sr0.font.size = Pt(10)
                    sr0.font.bold = True
                    sr0.font.color.rgb = RGBColor(255, 255, 255)
                    
                    step_c1 = flow_tbl.rows[s_idx].cells[1]
                    step_c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    set_cell_background(step_c1, 'F3F4F6' if s_idx % 2 == 1 else 'FFFFFF')
                    set_cell_margins(step_c1, top=60, bottom=60, left=100, right=80)
                    sp1 = step_c1.paragraphs[0]
                    sp1.paragraph_format.space_before = Pt(0)
                    sp1.paragraph_format.space_after = Pt(0)
                    step_txt = cells[s_idx][1].text if isinstance(cells[s_idx][1], Paragraph) else str(cells[s_idx][1])
                    add_formatted_runs(sp1, step_txt, base_font_size=10.0, base_color='111827')
                    
                set_table_borders(flow_tbl, color='E5E7EB', sz='4', val='single')
                
                p_sp = doc.add_paragraph()
                p_sp.paragraph_format.space_before = Pt(0)
                p_sp.paragraph_format.space_after = Pt(6)
                return

            # 6e. Real Data Tables
            word_tbl = doc.add_table(rows=r_count, cols=c_count)
            word_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            word_tbl.autofit = False
            
            col_widths = getattr(el, '_colWidths', None)
            total_w = 18.0
            if col_widths and len(col_widths) == c_count and all(w is not None for w in col_widths):
                sum_w = sum(col_widths)
                calculated_widths = [w / sum_w * total_w for w in col_widths]
            else:
                calculated_widths = [total_w / c_count] * c_count
                
            for ci, col in enumerate(word_tbl.columns):
                col.width = Cm(calculated_widths[ci])
                
            for ri, row in enumerate(cells):
                w_row = word_tbl.rows[ri]
                if ri == 0:
                    trPr = w_row._tr.get_or_add_trPr()
                    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
                trPr = w_row._tr.get_or_add_trPr()
                trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
                
                for ci, cell_val in enumerate(row):
                    w_cell = w_row.cells[ci]
                    w_cell.width = Cm(calculated_widths[ci])
                    w_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    
                    if ri == 0:
                        set_cell_background(w_cell, '1F2937')
                        txt_color = 'FFFFFF'
                        is_b = True
                    else:
                        bg_color = 'F3F4F6' if ri % 2 == 1 else 'FFFFFF'
                        set_cell_background(w_cell, bg_color)
                        txt_color = '111827'
                        is_b = False
                        
                    set_cell_margins(w_cell, top=60, bottom=60, left=60, right=60)
                    
                    p_cell = w_cell.paragraphs[0]
                    p_cell.paragraph_format.space_before = Pt(0)
                    p_cell.paragraph_format.space_after = Pt(0)
                    p_cell.paragraph_format.line_spacing = 1.05
                    
                    if ri == 0:
                        p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        p_cell.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        
                    cell_text = cell_val.text if isinstance(cell_val, Paragraph) else str(cell_val)
                    font_sz = 8.5 if c_count > 4 else 9.5
                    add_formatted_runs(p_cell, cell_text, base_font_size=font_sz, base_color=txt_color, default_bold=is_b)
                    
            set_data_table_borders(word_tbl, color='D1D5DB', sz='4', val='single')
            
            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.space_before = Pt(0)
            p_sp.paragraph_format.space_after = Pt(6)
            return

    # First pass: if title hasn't been emitted yet, look for title in story or create default
    for el in story:
        if not title_emitted:
            if isinstance(el, Table) and len(el._cellvalues) == 1 and len(el._cellvalues[0]) >= 2:
                c0 = el._cellvalues[0][0]
                c1 = el._cellvalues[0][1]
                if isinstance(c0, Drawing) and isinstance(c1, Paragraph) and 'Title' in c1.style.name:
                    process_story_element(el)
                    continue
            elif isinstance(el, Paragraph) and 'Title' in el.style.name:
                process_story_element(el)
                continue

    if not title_emitted:
        # Emit default title
        p_chap = doc.add_paragraph()
        p_chap.paragraph_format.space_before = Pt(0)
        p_chap.paragraph_format.space_after = Pt(2)
        r_chap = p_chap.add_run(f"CHAPTER {ch_num}")
        r_chap.font.name = "Times New Roman"
        r_chap.font.size = Pt(11)
        r_chap.font.bold = True
        r_chap.font.color.rgb = RGBColor(107, 114, 128)
        
        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_before = Pt(0)
        p_title.paragraph_format.space_after = Pt(6)
        r_title = p_title.add_run(chapter_title)
        r_title.font.name = "Times New Roman"
        r_title.font.size = Pt(24)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(17, 24, 39)
        title_emitted = True

    # Main processing loop
    for el in story:
        process_story_element(el)

    os.makedirs(os.path.dirname(out_docx_path), exist_ok=True)
    doc.save(out_docx_path)
    file_size_kb = os.path.getsize(out_docx_path) / 1024
    print(f"Generated DOCX ({file_size_kb:.1f} KB) -> {out_docx_path}")
    return out_docx_path

if __name__ == '__main__':
    all_target_chapters = [
        # Class 11 (Excluding Human Physiology unit: Ch14-Ch19)
        (1, 'The Living World', 'Ch1_TheLivingWorld', 'notes/class 11/Ch1_TheLivingWorld', 'notes/class 11/Ch1_TheLivingWorld/Ch1_TheLivingWorld.docx'),
        (2, 'Biological Classification', 'Ch2_BiologicalClassification', 'notes/class 11/Ch2_BiologicalClassification', 'notes/class 11/Ch2_BiologicalClassification/Ch2_BiologicalClassification.docx'),
        (3, 'Plant Kingdom', 'Ch3_PlantKingdom', 'notes/class 11/Ch3_PlantKingdom', 'notes/class 11/Ch3_PlantKingdom/Ch3_PlantKingdom.docx'),
        (4, 'Animal Kingdom', 'Ch4_AnimalKingdom', 'notes/class 11/Ch4_AnimalKingdom', 'notes/class 11/Ch4_AnimalKingdom/Ch4_AnimalKingdom.docx'),
        (5, 'Morphology of Flowering Plants', 'Ch5_MorphologyOfFloweringPlants', 'notes/class 11/Ch5_MorphologyOfFloweringPlants', 'notes/class 11/Ch5_MorphologyOfFloweringPlants/Ch5_MorphologyOfFloweringPlants.docx'),
        (6, 'Anatomy of Flowering Plants', 'Ch6_AnatomyOfFloweringPlants', 'notes/class 11/Ch6_AnatomyOfFloweringPlants', 'notes/class 11/Ch6_AnatomyOfFloweringPlants/Ch6_AnatomyOfFloweringPlants.docx'),
        (7, 'Structural Organisation in Animals', 'Ch7_StructuralOrganisationInAnimals', 'notes/class 11/Ch7_StructuralOrganisationInAnimals', 'notes/class 11/Ch7_StructuralOrganisationInAnimals/Ch7_StructuralOrganisationInAnimals.docx'),
        (8, 'Cell: The Unit of Life', 'Ch8_CellTheUnitOfLife', 'notes/class 11/Ch8_CellTheUnitOfLife', 'notes/class 11/Ch8_CellTheUnitOfLife/Ch8_CellTheUnitOfLife.docx'),
        (9, 'Biomolecules', 'Ch9_Biomolecules', 'notes/class 11/Ch9_Biomolecules', 'notes/class 11/Ch9_Biomolecules/Ch9_Biomolecules.docx'),
        (10, 'Cell Cycle and Cell Division', 'Ch10_CellCycleAndCellDivision', 'notes/class 11/Ch10_CellCycleAndCellDivision', 'notes/class 11/Ch10_CellCycleAndCellDivision/Ch10_CellCycleAndCellDivision.docx'),
        (11, 'Photosynthesis in Higher Plants', 'Ch11_PhotosynthesisInHigherPlants', 'notes/class 11/Ch11_PhotosynthesisInHigherPlants', 'notes/class 11/Ch11_PhotosynthesisInHigherPlants/Ch11_PhotosynthesisInHigherPlants.docx'),
        (12, 'Respiration in Plants', 'Ch12_RespirationInPlants', 'notes/class 11/Ch12_RespirationInPlants', 'notes/class 11/Ch12_RespirationInPlants/Ch12_RespirationInPlants.docx'),
        (13, 'Plant Growth and Development', 'Ch13_PlantGrowthAndDevelopment', 'notes/class 11/Ch13_PlantGrowthAndDevelopment', 'notes/class 11/Ch13_PlantGrowthAndDevelopment/Ch13_PlantGrowthAndDevelopment.docx'),
        
        # Class 12 (Excluding Ch5 Molecular Basis of Inheritance and Ch6 Evolution)
        (1, 'Sexual Reproduction in Flowering Plants', 'Ch1_SexualReproductionInFloweringPlants', 'notes/class 12/Ch1_SexualReproductionInFloweringPlants', 'notes/class 12/Ch1_SexualReproductionInFloweringPlants/Ch1_SexualReproductionInFloweringPlants.docx'),
        (2, 'Human Reproduction', 'Ch2_HumanReproduction', 'notes/class 12/Ch2_HumanReproduction', 'notes/class 12/Ch2_HumanReproduction/Ch2_HumanReproduction.docx'),
        (3, 'Reproductive Health', 'Ch3_ReproductiveHealth', 'notes/class 12/Ch3_ReproductiveHealth', 'notes/class 12/Ch3_ReproductiveHealth/Ch3_ReproductiveHealth.docx'),
        (4, 'Principles of Inheritance and Variation', 'Ch4_PrinciplesOfInheritanceAndVariation', 'notes/class 12/Ch4_PrinciplesOfInheritanceAndVariation', 'notes/class 12/Ch4_PrinciplesOfInheritanceAndVariation/Ch4_PrinciplesOfInheritanceAndVariation.docx'),
        (7, 'Human Health and Disease', 'Ch7_HumanHealthAndDisease', 'notes/class 12/Ch7_HumanHealthAndDisease', 'notes/class 12/Ch7_HumanHealthAndDisease/Ch7_HumanHealthAndDisease.docx'),
        (8, 'Microbes in Human Welfare', 'Ch8_MicrobesInHumanWelfare', 'notes/class 12/Ch8_MicrobesInHumanWelfare', 'notes/class 12/Ch8_MicrobesInHumanWelfare/Ch8_MicrobesInHumanWelfare.docx'),
        (9, 'Biotechnology: Principles and Processes', 'Ch9_BiotechnologyPrinciplesAndProcesses', 'notes/class 12/Ch9_BiotechnologyPrinciplesAndProcesses', 'notes/class 12/Ch9_BiotechnologyPrinciplesAndProcesses/Ch9_BiotechnologyPrinciplesAndProcesses.docx'),
        (10, 'Biotechnology and its Applications', 'Ch10_BiotechnologyAndItsApplications', 'notes/class 12/Ch10_BiotechnologyAndItsApplications', 'notes/class 12/Ch10_BiotechnologyAndItsApplications/Ch10_BiotechnologyAndItsApplications.docx'),
        (11, 'Organisms and Populations', 'Ch11_OrganismsAndPopulations', 'notes/class 12/Ch11_OrganismsAndPopulations', 'notes/class 12/Ch11_OrganismsAndPopulations/Ch11_OrganismsAndPopulations.docx'),
        (12, 'Ecosystem', 'Ch12_Ecosystem', 'notes/class 12/Ch12_Ecosystem', 'notes/class 12/Ch12_Ecosystem/Ch12_Ecosystem.docx'),
        (13, 'Biodiversity and Conservation', 'Ch13_BiodiversityAndConservation', 'notes/class 12/Ch13_BiodiversityAndConservation', 'notes/class 12/Ch13_BiodiversityAndConservation/Ch13_BiodiversityAndConservation.docx'),
    ]
    
    print(f"Beginning DOCX generation for {len(all_target_chapters)} chapters...")
    successes = []
    failures = []
    
    for ch_num, ch_title, mod_name, ch_dir, out_docx in all_target_chapters:
        try:
            build_docx_for_chapter(ch_num, ch_title, mod_name, ch_dir, out_docx)
            successes.append((ch_num, ch_title, out_docx))
        except Exception as e:
            print(f"ERROR generating {mod_name}: {e}")
            import traceback
            traceback.print_exc()
            failures.append((ch_num, ch_title, str(e)))
            
    print("\n" + "="*60)
    print(f"SUMMARY: {len(successes)} succeeded, {len(failures)} failed.")
    print("="*60)
    if failures:
        print("Failed chapters:")
        for ch_num, ch_title, err in failures:
            print(f"  - Ch {ch_num} ({ch_title}): {err}")
